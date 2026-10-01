import openpyxl
import os
from collections import defaultdict

def analyze_dataset(file_path):
    print("Loading workbook...")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    sheet_names = wb.sheetnames
    num_sheets = len(sheet_names)
    
    print(f"Workbook sheets: {sheet_names}")
    
    sheet_info = {}
    for name in sheet_names:
        sheet = wb[name]
        sheet_info[name] = {
            'rows': sheet.max_row,
            'cols': sheet.max_column
        }
    
    # Identify transactions sheet
    txn_sheet_name = None
    if "Transactions" in sheet_names:
        txn_sheet_name = "Transactions"
    else:
        for name in sheet_names:
            # check headers
            sheet = wb[name]
            headers = [str(cell.value) for cell in sheet[1] if cell.value is not None]
            if any(h in headers for h in ["Transaction_ID", "Amount", "Sender_Account"]):
                txn_sheet_name = name
                break
                
    if not txn_sheet_name:
        print("No transactions sheet found.")
        return
        
    sheet = wb[txn_sheet_name]
    headers = [str(cell.value) for cell in sheet[1]]
    total_txn_rows = sheet.max_row - 1
    
    print(f"Transactions sheet: {txn_sheet_name}")
    print(f"Headers: {headers}")
    
    expected_fields = [
        "Transaction_ID", "Sender_Account", "Receiver_Account", 
        "Sender_IFSC", "Receiver_IFSC", "Amount", "Timestamp", 
        "Payment_Mode", "Narration", "IP_Address", "Device_Type"
    ]
    
    present_fields = [h for h in headers if h in expected_fields]
    missing_fields = [h for h in expected_fields if h not in headers]
    unexpected_fields = [h for h in headers if h not in expected_fields and h != 'None']
    
    print(f"Present fields: {present_fields}")
    print(f"Missing fields: {missing_fields}")
    print(f"Unexpected fields: {unexpected_fields}")
    
    # Need to inspect the underlying types and values
    col_idx = {h: i for i, h in enumerate(headers)}
    
    # Data to collect
    types_found = defaultdict(set)
    
    tx_ids = []
    missing_tx = 0
    duplicate_tx = 0
    
    senders = set()
    missing_senders = 0
    
    receivers = set()
    missing_receivers = 0
    
    sender_receiver_pairs = []
    
    amounts = []
    missing_amounts = 0
    zero_amounts = 0
    negative_amounts = 0
    invalid_amounts = 0
    
    timestamps = []
    missing_timestamps = 0
    invalid_timestamps = 0
    
    payment_modes = defaultdict(int)
    
    narrations = set()
    missing_narrations = 0
    first_20_narrations = []
    
    ips = set()
    missing_ips = 0
    first_10_ips = []
    
    devices = defaultdict(int)
    missing_devices = 0
    
    first_5 = []
    last_5 = []
    
    # Start loop
    count = 0
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        count += 1
        
        # Save first 5 and last 5
        row_vals = [cell.value for cell in row]
        if count <= 5:
            first_5.append(row_vals)
        if count > total_txn_rows - 5:
            last_5.append(row_vals)
            
        for i, cell in enumerate(row):
            if i < len(headers):
                h = headers[i]
                if cell.value is not None:
                    types_found[h].add(type(cell.value).__name__)
        
        # Tx ID
        if "Transaction_ID" in col_idx:
            tx_id = row[col_idx["Transaction_ID"]].value
            if tx_id is None:
                missing_tx += 1
            else:
                tx_ids.append(str(tx_id))
                
        # Senders and Receivers
        sender = None
        receiver = None
        if "Sender_Account" in col_idx:
            sender = row[col_idx["Sender_Account"]].value
            if sender is None:
                missing_senders += 1
            else:
                sender = str(sender)
                senders.add(sender)
                
        if "Receiver_Account" in col_idx:
            receiver = row[col_idx["Receiver_Account"]].value
            if receiver is None:
                missing_receivers += 1
            else:
                receiver = str(receiver)
                receivers.add(receiver)
                
        if sender and receiver:
            sender_receiver_pairs.append((sender, receiver))
            
        # Amount
        if "Amount" in col_idx:
            amt = row[col_idx["Amount"]].value
            if amt is None:
                missing_amounts += 1
            else:
                try:
                    f_amt = float(amt)
                    amounts.append(f_amt)
                    if f_amt == 0:
                        zero_amounts += 1
                    elif f_amt < 0:
                        negative_amounts += 1
                except ValueError:
                    invalid_amounts += 1
                    
        # Timestamp
        if "Timestamp" in col_idx:
            ts = row[col_idx["Timestamp"]].value
            if ts is None:
                missing_timestamps += 1
            else:
                # Can be datetime or string
                timestamps.append(ts)
                
        # Payment Mode
        if "Payment_Mode" in col_idx:
            pm = row[col_idx["Payment_Mode"]].value
            if pm is not None:
                payment_modes[str(pm)] += 1
                
        # Narration
        if "Narration" in col_idx:
            nar = row[col_idx["Narration"]].value
            if nar is None:
                missing_narrations += 1
            else:
                s_nar = str(nar)
                narrations.add(s_nar)
                if len(first_20_narrations) < 20:
                    first_20_narrations.append(s_nar)
                    
        # IP
        if "IP_Address" in col_idx:
            ip = row[col_idx["IP_Address"]].value
            if ip is None:
                missing_ips += 1
            else:
                s_ip = str(ip)
                ips.add(s_ip)
                if len(first_10_ips) < 10:
                    first_10_ips.append(s_ip)
                    
        # Device
        if "Device_Type" in col_idx:
            dev = row[col_idx["Device_Type"]].value
            if dev is None:
                missing_devices += 1
            else:
                devices[str(dev)] += 1
                
    # Tx IDs
    total_tx = len(tx_ids)
    unique_tx = len(set(tx_ids))
    duplicate_tx = total_tx - unique_tx
    
    unique_senders = len(senders)
    unique_receivers = len(receivers)
    all_accounts = senders.union(receivers)
    both_accounts = senders.intersection(receivers)
    
    unique_pairs = len(set(sender_receiver_pairs))
    repeated_pairs = len(sender_receiver_pairs) - unique_pairs
    
    if amounts:
        total_amt = sum(amounts)
        min_amt = min(amounts)
        max_amt = max(amounts)
        avg_amt = total_amt / len(amounts)
    else:
        total_amt = min_amt = max_amt = avg_amt = 0
        
    # Generate Report
    report = []
    report.append("============================================================")
    report.append("TraceX — Phase 1 Dataset Validation Report")
    report.append("============================================================")
    report.append("")
    report.append("1. Product Information")
    report.append("Name: TraceX")
    report.append("Type: Cyber Money Trail Intelligence Platform")
    report.append("")
    report.append("2. Requirements Source")
    report.append("Source: Problem Statement.pdf")
    report.append("")
    report.append("3. Dataset Information")
    report.append("Source: data/Dataset_VoidHacks.xlsx")
    report.append("")
    report.append("4. Workbook Information")
    report.append(f"Number of sheets: {num_sheets}")
    report.append(f"Sheet names: {sheet_names}")
    report.append("")
    report.append("5. Sheet Information")
    for name, info in sheet_info.items():
        report.append(f"Sheet '{name}': {info['rows']} rows, {info['cols']} columns")
    report.append("")
    report.append("6. Transaction Row Count")
    report.append(f"Total transaction rows: {total_txn_rows}")
    report.append("First 5 records:")
    for r in first_5: report.append(str(r))
    report.append("Last 5 records:")
    for r in last_5: report.append(str(r))
    report.append("")
    report.append("7. Expected Dataset Schema")
    report.append(", ".join(expected_fields))
    report.append("")
    report.append("8. Actual Dataset Schema")
    report.append(", ".join(headers))
    report.append("")
    report.append("9. Present Fields")
    report.append(", ".join(present_fields))
    report.append("")
    report.append("10. Missing Required Fields")
    if missing_fields:
        for mf in missing_fields:
            report.append(f"{mf} is specified by the Problem Statement but is not present in the supplied Excel dataset.")
    else:
        report.append("None")
    report.append("")
    report.append("11. Unexpected Fields")
    if unexpected_fields:
        report.append(", ".join(unexpected_fields))
    else:
        report.append("None")
    report.append("")
    report.append("12. Data Types")
    for field, types in types_found.items():
        report.append(f"{field}: {', '.join(types)}")
    report.append("")
    report.append("13. Transaction ID Quality")
    report.append(f"Total: {total_tx}")
    report.append(f"Missing: {missing_tx}")
    report.append(f"Unique: {unique_tx}")
    report.append(f"Duplicate: {duplicate_tx}")
    report.append("First 10 examples:")
    report.append(", ".join(tx_ids[:10]))
    report.append("")
    report.append("14. Account Quality")
    report.append(f"Missing Sender Accounts: {missing_senders}")
    report.append(f"Missing Receiver Accounts: {missing_receivers}")
    report.append(f"Unique Sender Accounts: {unique_senders}")
    report.append(f"Unique Receiver Accounts: {unique_receivers}")
    report.append(f"Unique Accounts Overall: {len(all_accounts)}")
    report.append(f"Accounts appearing as both: {len(both_accounts)}")
    report.append("")
    report.append("15. Sender/Receiver Relationship Statistics")
    report.append(f"Unique Sender -> Receiver pairs: {unique_pairs}")
    report.append(f"Repeated Sender -> Receiver pairs: {repeated_pairs}")
    report.append("")
    report.append("16. Amount Statistics")
    report.append(f"Total Amount: {total_amt}")
    report.append(f"Minimum Amount: {min_amt}")
    report.append(f"Maximum Amount: {max_amt}")
    report.append(f"Average Amount: {avg_amt}")
    report.append(f"Zero Amount Count: {zero_amounts}")
    report.append(f"Negative Amount Count: {negative_amounts}")
    report.append(f"Missing Amount Count: {missing_amounts}")
    report.append(f"Invalid Amount Count: {invalid_amounts}")
    report.append("")
    report.append("17. Timestamp Statistics")
    # check valid vs invalid
    import datetime
    valid_dt = 0
    invalid_dt = 0
    earliest = None
    latest = None
    for ts in timestamps:
        if isinstance(ts, datetime.datetime):
            valid_dt += 1
            if earliest is None or ts < earliest: earliest = ts
            if latest is None or ts > latest: latest = ts
        else:
            try:
                # just attempt basic parsing if it's string
                from dateutil import parser
                dt = parser.parse(str(ts))
                valid_dt += 1
                if earliest is None or dt < earliest: earliest = dt
                if latest is None or dt > latest: latest = dt
            except:
                invalid_dt += 1
    report.append(f"Underlying Type: {[t for t in types_found.get('Timestamp', [])]}")
    report.append(f"Valid Datetime Count: {valid_dt}")
    report.append(f"Invalid Datetime Count: {invalid_dt}")
    report.append(f"Missing Datetime Count: {missing_timestamps}")
    report.append(f"Earliest Valid Timestamp: {earliest}")
    report.append(f"Latest Valid Timestamp: {latest}")
    if earliest and latest:
        report.append(f"Total Time Span: {latest - earliest}")
    report.append("Sample Timestamps: " + str(timestamps[:5]))
    report.append("")
    report.append("18. Payment Mode Statistics")
    for mode, c in payment_modes.items():
        report.append(f"{mode} | {c}")
    report.append("")
    report.append("19. Narration Statistics")
    report.append(f"Missing Count: {missing_narrations}")
    report.append(f"Unique Count: {len(narrations)}")
    report.append("First 20 examples:")
    for n in first_20_narrations: report.append(n)
    report.append("")
    report.append("20. IP Address Statistics")
    report.append(f"Missing Count: {missing_ips}")
    report.append(f"Unique Count: {len(ips)}")
    report.append("First 10 examples:")
    report.append(", ".join(first_10_ips))
    report.append("")
    report.append("21. Device Type Statistics")
    if "Device_Type" in col_idx:
        report.append(f"Missing Count: {missing_devices}")
        report.append(f"Unique Values: {len(devices)}")
        for dev, c in devices.items():
            report.append(f"{dev}: {c}")
    else:
        report.append("Device_Type is specified by the Problem Statement but is not present in the supplied Excel dataset.")
    report.append("")
    report.append("22. Summary Sheet Information")
    if "Summary" in sheet_names:
        report.append("Summary sheet present.")
        # could compare but for now just acknowledge
    else:
        report.append("No Summary sheet present.")
    report.append("")
    report.append("23. Data Quality Issues")
    issues = []
    if duplicate_tx > 0: issues.append("Duplicate Transaction IDs found.")
    if missing_tx > 0: issues.append("Missing Transaction IDs found.")
    if missing_senders > 0: issues.append("Missing Sender Accounts found.")
    if missing_receivers > 0: issues.append("Missing Receiver Accounts found.")
    if missing_amounts > 0: issues.append("Missing Amounts found.")
    if invalid_amounts > 0: issues.append("Invalid Amounts found.")
    if zero_amounts > 0: issues.append("Zero Amounts found.")
    if negative_amounts > 0: issues.append("Negative Amounts found.")
    if missing_timestamps > 0: issues.append("Missing Timestamps found.")
    if missing_fields: issues.append("Schema mismatch: Missing fields.")
    report.append("\n".join(issues) if issues else "No major issues identified.")
    report.append("")
    report.append("24. Dataset Scale Comparison")
    report.append("Required dataset scale: 2,000,000 Banking Records")
    report.append(f"Actual supplied dataset: {total_txn_rows} records")
    report.append("")
    report.append("25. Limitations")
    report.append("This dataset scale is limited compared to the stated requirement. Some fields may be absent.")
    report.append("")
    report.append("26. Recommendations for Phase 2")
    report.append("Proceed with Module B (Mule Ring Detection) using the actual available dataset schema.")
    
    with open("reports/TraceX_Phase1_Dataset_Report.txt", "w") as f:
        f.write("\n".join(report))
        
    print("Report generated at reports/TraceX_Phase1_Dataset_Report.txt")

if __name__ == "__main__":
    analyze_dataset("data/Dataset_VoidHacks.xlsx")
